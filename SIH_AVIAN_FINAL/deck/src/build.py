from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

BLUE = RGBColor(0x00, 0x70, 0xC0)      # template footer blue — dominant accent
NAVY = RGBColor(0x1F, 0x2A, 0x44)      # body ink
MUTED = RGBColor(0x59, 0x62, 0x70)
TINT = RGBColor(0xEA, 0xF2, 0xFB)      # light blue card fill
AMBER = RGBColor(0xB4, 0x5F, 0x06)     # risk accent
AMBER_TINT = RGBColor(0xFD, 0xF1, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

TEAM_NAME = "TRINETRA"
TEAM_ID = "129300"

prs = Presentation("template.pptx")


# ---------- helpers ----------
def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tb


def add_para(tf, runs, size=12, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
             space_after=0, bullet=False, first=False, level=0):
    """runs: str or list of (text, dict(bold/color/size/link/italic))"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    if space_after:
        p.space_after = Pt(space_after)
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", False)
        f.color.rgb = opt.get("color", color)
        if opt.get("underline"):
            f.underline = True
        if opt.get("link"):
            r.hyperlink.address = opt["link"]
    if bullet:
        _bullet(p)
    return p


def _bullet(p):
    from pptx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.18))))
    pPr.set("indent", str(-Emu(Inches(0.18))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    from lxml import etree
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def card(slide, x, y, w, h, fill=TINT, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pointer(slide, x, y, w, text, size=15, color=BLUE):
    """A template 'idea details pointer', kept word-for-word, used as a heading."""
    tb = textbox(slide, x, y, w, 0.35)
    add_para(tb.text_frame, text, size=size, bold=True, color=color, first=True)
    return tb


def set_title(slide, text, size=28):
    t = shape_by_name(slide, "Title 1")
    tf = t.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    # keep the first (empty, carries the vertical-tab) run's formatting, set text on the second
    for r in runs:
        r.text = ""
    target = runs[-1]
    target.text = text
    for r in runs:
        r.font.size = Pt(size)


def set_team(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = TEAM_NAME
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.bold = True
            for r in p.runs[1:]:
                r.text = ""
            sh.text_frame.word_wrap = False


def picture_cropped(slide, path, x, y, w, h):
    im = Image.open(path)
    iw, ih = im.size
    target = w / h
    src = iw / ih
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if src > target:  # too wide -> crop sides
        crop = (1 - target / src) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - src / target) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


# ---------- drop the instructions slide (7) ----------
sldIdLst = prs.slides._sldIdLst
last = sldIdLst[-1]
prs.part.drop_rel(last.rId)
sldIdLst.remove(last)

s1, s2, s3, s4, s5, s6 = prs.slides

# =========================================================
# SLIDE 1 — Title page (fill the six template lines)
# =========================================================
tb = shape_by_name(s1, "TextBox 9")
values = {
    "Problem Statement ID –": " SIH26201",
    "Problem Statement Title-": (" Student Innovation-There is a need to design drones and robots "
                                 "that can solve some of the pressing challenges of India such as "
                                 "handling medical emergencies, search and rescue operations, etc."),
    "Theme-": " Robotics and Drones",
    "PS Category- Software/Hardware": None,  # replaced below
    "Team ID-": " " + TEAM_ID,
    "Team Name (Registered on portal)": " - @TRINETRA",
}
for p in tb.text_frame.paragraphs:
    if not p.runs:
        continue
    label = p.runs[0].text
    p.runs[0].font.size = Pt(15)
    if label == "PS Category- Software/Hardware":
        p.runs[0].text = "PS Category-"
        val = " Software"
    else:
        val = values.get(label)
    if val is None:
        continue
    r = p.add_run()
    r.text = val
    r.font.name = FONT
    r.font.bold = False
    r.font.size = Pt(12 if label.startswith("Problem Statement Title") else 15)
    r.font.color.rgb = NAVY
tb.text_frame.word_wrap = True

# =========================================================
# SLIDE 2 — Idea title + proposed solution
# =========================================================
set_team(s2)
set_title(s2, "AVIAN: Autonomous Bridge-Inspection Drone", size=28)
remove_shape(shape_by_name(s2, "TextBox 8"))

tb = textbox(s2, 0.5, 1.28, 12.3, 0.4)
add_para(tb.text_frame,
         [("Proposed Solution (Describe your Idea/Solution/Prototype)",
           {"underline": True})],
         size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79), first=True)

LX, LW = 0.5, 7.35
blocks = [
    ("Detailed explanation of the proposed solution", [
        "Plans its own inspection route from the bridge's 3D model — no list of known defects needed",
        "Flies the route autonomously, steering clear of steel and concrete using on-board range sensing",
        "A zoom camera captures millimetre-level detail from a safe 8 m standoff",
        "AI flags cracks, spalling and corrosion and tags each one to the exact bridge member",
    ]),
    ("How it addresses the problem", [
        "India's railways alone have 1,47,523 bridges; 37,689 are over 100 years old. They are checked "
        "“by visual perception” with ~40% of bridge-staff posts vacant",
        "Every span gets a repeatable, geo-tagged image record instead of one person's judgement",
        "Reaches under-deck and truss zones with no rope access, scaffolding or lane closure",
    ]),
    ("Innovation and uniqueness of the solution", [
        "Route planned from geometry only — code-audited to read zero defect data",
        "“Survey + zoom” instead of flying close: keeps ≥ 3 m from structure, as China's 2026 "
        "UAV bridge-inspection guideline requires",
        "Every run is scored against a digital twin with 192 known defects, so results are measured, not claimed",
    ]),
]
y = 1.82
heights = [1.5, 1.62, 1.6]
for (label, bullets), h in zip(blocks, heights):
    tb = textbox(s2, LX, y, LW, h)
    tf = tb.text_frame
    add_para(tf, label, size=14, bold=True, color=BLUE, first=True, space_after=3)
    for b in bullets:
        add_para(tf, b, size=11.5, bullet=True, space_after=2)
    y += h

# right column: the twin + the local collapse that motivates it
RX, RW = 8.15, 4.68
picture_cropped(s2, "CAM_10_TRUSS.jpg", RX, 1.85, RW, 2.63)
tb = textbox(s2, RX, 4.53, RW, 0.35)
add_para(tb.text_frame,
         "Our digital twin: 360 m road bridge + steel truss + metro viaduct, with 192 measured defects",
         size=9.5, color=MUTED, first=True)

c = card(s2, RX, 5.02, RW, 1.62, fill=AMBER_TINT)
tb = textbox(s2, RX + 0.2, 5.14, RW - 0.4, 1.45)
tf = tb.text_frame
add_para(tf, "Why now — 9 July 2025", size=13, bold=True, color=AMBER, first=True, space_after=3)
add_para(tf, "The 40-year-old Gambhira bridge linking Vadodara and Anand districts collapsed "
             "into the Mahisagar river, killing 22. Residents had reportedly warned about its condition "
             "for years; four engineers were suspended for lapses in maintenance and supervision. (ref. 2)",
         size=11, color=NAVY)

# =========================================================
# SLIDE 3 — Technical approach
# =========================================================
set_team(s3)
remove_shape(shape_by_name(s3, "TextBox 8"))
pointer(s3, 0.5, 1.28, 12.3,
        "Technologies to be used (e.g. programming languages, frameworks, hardware)")

cards = [
    ("Digital twin & simulation", [
        "Blender: bridge model, 192 measured defects",
        "PyBullet: flight physics",
        "Gazebo + ROS 2 Jazzy: exported world",
    ]),
    ("Autonomy (Python)", [
        "Cascade PID flight controller",
        "Coverage planner (greedy set-cover)",
        "Ray-cast range sensor for avoidance",
    ]),
    ("AI defect detection", [
        "PyTorch + torchvision Faster R-CNN (MobileNetV3)",
        "BSD-3 licence — free to hand over",
        "Trained on rendered defect images",
    ]),
    ("Target hardware (planned)", [
        "PX4 quadcopter + zoom gimbal camera",
        "LiDAR / visual-inertial SLAM",
        "Jetson-class on-board computer",
    ]),
]
cw, gap = 2.94, 0.19
for i, (head, lines) in enumerate(cards):
    x = 0.5 + i * (cw + gap)
    card(s3, x, 1.75, cw, 1.62)
    tb = textbox(s3, x + 0.16, 1.86, cw - 0.32, 1.45)
    tf = tb.text_frame
    add_para(tf, head, size=13, bold=True, color=BLUE, first=True, space_after=4)
    for ln in lines:
        add_para(tf, ln, size=11, bullet=True, space_after=2)

pointer(s3, 0.5, 3.62, 12.3,
        "Methodology and process for implementation (Flow Charts/Images/ working prototype)")

steps = [
    ("1", "Digital twin", "Bridge 3D model → 568 collision shapes"),
    ("2", "Plan coverage", "150 viewpoints from geometry only"),
    ("3", "Fly & avoid", "Ray-cast sensing steers clear of structure"),
    ("4", "Zoom capture", "~9° zoom view from 8 m, at the pose actually reached"),
    ("5", "Detect", "Faster R-CNN flags defects in each image"),
    ("6", "Score", "Checked against 192 known defects; report per bridge member"),
]
bw, bgap = 1.8, 0.306
by, bh = 4.08, 1.3
for i, (num, head, detail) in enumerate(steps):
    x = 0.5 + i * (bw + bgap)
    box = card(s3, x, by, bw, bh, fill=WHITE, line=BLUE)
    circ = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12), Inches(by + 0.12),
                               Inches(0.36), Inches(0.36))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE
    circ.line.fill.background(); circ.shadow.inherit = False
    ctf = circ.text_frame
    ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
    add_para(ctf, num, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tb = textbox(s3, x + 0.56, by + 0.15, bw - 0.66, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, head, size=11.5, bold=True, color=NAVY, first=True)
    tb = textbox(s3, x + 0.12, by + 0.6, bw - 0.24, bh - 0.68)
    add_para(tb.text_frame, detail, size=10.5, color=MUTED, first=True)
    if i < len(steps) - 1:
        ax = x + bw + 0.05
        arr = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(by + bh / 2 - 0.11),
                                  Inches(bgap - 0.1), Inches(0.22))
        arr.fill.solid(); arr.fill.fore_color.rgb = BLUE
        arr.line.fill.background(); arr.shadow.inherit = False

card(s3, 0.5, 5.7, 12.33, 0.95)
tb = textbox(s3, 0.72, 5.78, 11.9, 0.8, anchor=MSO_ANCHOR.MIDDLE)
tf = tb.text_frame
add_para(tf, [("Working prototype today: ", {"bold": True, "color": BLUE}),
              ("all six steps run end-to-end in simulation, with measured results on the next slide. ", {}),
              ("Next: ", {"bold": True, "color": BLUE}),
              ("run the same stack live in Gazebo + ROS 2, then on PX4 hardware.", {})],
         size=12.5, first=True)

# =========================================================
# SLIDE 4 — Feasibility and viability
# =========================================================
set_team(s4)
remove_shape(shape_by_name(s4, "TextBox 8"))

pointer(s4, 0.5, 1.28, 5.4, "Analysis of the feasibility of the idea")
tb = textbox(s4, 0.5, 1.62, 5.4, 0.3)
add_para(tb.text_frame, "Measured in our simulation — not estimated", size=10.5,
         color=MUTED, first=True)

stats = [
    ("56.7%", "of bridge surfaces seen by the camera — route planned with no defect data (150 waypoints)"),
    ("39 / 73", "known defects fell inside a settled camera view (53.4%)"),
    ("239 / 239", "flight-log entries show avoidance driven by the sensor, not a map (6,492 reactions)"),
    ("71% → 45%", "waypoints with a collision or stall, before vs after sensed avoidance"),
    ("0.34", "detector mAP@0.5 on 24 held-out defects (synthetic images)"),
]
sy = 2.02
for big, label in stats:
    tb = textbox(s4, 0.5, sy, 1.75, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, big, size=20, bold=True, color=BLUE, first=True)
    tb = textbox(s4, 2.35, sy, 3.55, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, label, size=11, color=NAVY, first=True)
    sy += 0.8

# right: risks -> strategies table (template pointers as column headers)
TX, TW = 6.2, 6.63
rows = [
    ("Detection accuracy is the weakest link: only 1–2 of 10 defects found in zoomed mission images "
     "(two model versions); "
     "loose bolts are too small at current camera resolution",
     "Retrain on real bridge datasets (CODEBRIM, dacl10k) on a GPU; higher-resolution capture for "
     "bolts; an engineer confirms every flag before it is reported"),
    ("Collisions in the dense steel truss: 45% of waypoints still hit or stall in simulation",
     "Clearance checks along the whole path, not just at waypoints; slower close-range legs; "
     "target under 5% before any field flight"),
    ("No GPS under the deck; magnetic noise near 25 kV railway overhead lines",
     "Visual-inertial / LiDAR SLAM, as used by Flyability's Elios 3; compass-free heading"),
]
tbl_shape = s4.shapes.add_table(len(rows) + 1, 2, Inches(TX), Inches(1.32), Inches(TW), Inches(4.5))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.2)
tbl.columns[1].width = Inches(TW - 3.2)
tbl.first_row = True
heads = ["Potential challenges and risks", "Strategies for overcoming these challenges"]
for c, h in enumerate(heads):
    cell = tbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
    tf = cell.text_frame
    tf.paragraphs[0].text = ""
    add_para(tf, h, size=12, bold=True, color=WHITE, first=True)
    cell.margin_left = cell.margin_right = Inches(0.1)
    cell.margin_top = cell.margin_bottom = Inches(0.06)
tbl.rows[0].height = Inches(0.5)
for r, (risk, fix) in enumerate(rows, start=1):
    tbl.rows[r].height = Inches(1.3)
    for c, (txt, fill, col) in enumerate([(risk, AMBER_TINT, NAVY), (fix, WHITE, NAVY)]):
        cell = tbl.cell(r, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.06)
        tf = cell.text_frame
        tf.paragraphs[0].text = ""
        add_para(tf, txt, size=11, color=col, first=True)

card(s4, 0.5, 6.1, 12.33, 0.7)
tb = textbox(s4, 0.72, 6.13, 11.9, 0.64, anchor=MSO_ANCHOR.MIDDLE)
add_para(tb.text_frame,
         [("Regulatory path: ", {"bold": True, "color": BLUE}),
          ("Drone Rules 2021 R&D exemption at prototype stage → DGCA type certificate through a "
           "QCI-approved body (~₹1.5 lakh at NTH) → railway safety permits for any work near live 25 kV lines.",
           {})],
         size=11.5, first=True)

# =========================================================
# SLIDE 5 — Impact and benefits
# =========================================================
set_team(s5)
remove_shape(shape_by_name(s5, "TextBox 8"))

callouts = [
    ("1,47,523", "railway bridges in India"),
    ("37,689", "of them over 100 years old"),
    ("≈ 33", "bridges per bridge-staff member in post (1,47,523 ÷ 4,517)"),
]
cw5, g5 = 3.95, 0.24
for i, (big, label) in enumerate(callouts):
    x = 0.5 + i * (cw5 + g5)
    card(s5, x, 1.3, cw5, 1.25)
    tb = textbox(s5, x + 0.2, 1.36, cw5 - 0.4, 0.62, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, big, size=30, bold=True, color=BLUE, first=True)
    tb = textbox(s5, x + 0.2, 1.98, cw5 - 0.4, 0.5)
    add_para(tb.text_frame, label, size=11.5, color=NAVY, first=True)
tb = textbox(s5, 0.5, 2.6, 12.33, 0.25)
add_para(tb.text_frame, "Source: Standing Committee on Railways, Maintenance of Bridges in Indian Railways "
                        "(Jan 2019), ref. 1. Inspections rely on the “visual perception” of the inspecting official.",
         size=9.5, color=MUTED, first=True)

pointer(s5, 0.5, 3.0, 5.9, "Potential impact on the target audience")
tb = textbox(s5, 0.5, 3.42, 5.9, 2.45)
tf = tb.text_frame
aud = [
    ("Indian Railways — ", "fits the Railway Board's drone-inspection circulars and its Bridge Management System (4D BrIM)"),
    ("NHAI, MoRTH and state PWDs — ", "e.g. Gujarat's Roads & Buildings department after the Gambhira collapse"),
    ("Metro operators (DMRC) — ", "viaduct inspection without closing lines"),
    ("Rescue teams — ", "a quick structural check of a flood- or quake-damaged bridge before convoys cross"),
]
for i, (who, what) in enumerate(aud):
    add_para(tf, [(who, {"bold": True}), (what, {})], size=11.5, bullet=True,
             first=(i == 0), space_after=5)

tb = textbox(s5, 6.83, 3.0, 6.0, 0.6)
add_para(tb.text_frame, "Benefits of the solution (social, economic, environmental, etc.)", size=14, bold=True, color=BLUE, first=True)
ben = [
    ("Social", "Earlier warning of failing spans; inspectors off ropes and scaffolds"),
    ("Economic", "No lane closures or snooper trucks; each year's images map onto the same twin, so every defect is tracked over time"),
    ("Environmental", "Battery-electric; fewer heavy access vehicles and fewer emergency rebuilds"),
]
yb = 3.5
for head, text in ben:
    card(s5, 6.83, yb, 6.0, 0.68)
    tb = textbox(s5, 7.0, yb, 1.55, 0.68, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, head, size=12, bold=True, color=BLUE, first=True)
    tb = textbox(s5, 8.55, yb, 4.15, 0.68, anchor=MSO_ANCHOR.MIDDLE)
    add_para(tb.text_frame, text, size=11, color=NAVY, first=True)
    yb += 0.77

card(s5, 0.5, 6.05, 12.33, 0.72, fill=AMBER_TINT)
tb = textbox(s5, 0.72, 6.08, 11.9, 0.66, anchor=MSO_ANCHOR.MIDDLE)
add_para(tb.text_frame,
         [("Where it fits: ", {"bold": True, "color": AMBER}),
          ("between satellite radar, which tracks millimetre movement across whole regions but can't see "
           "under a deck, and manual close inspection, which can but is slow, risky and subjective.", {})],
         size=11.5, first=True)

# =========================================================
# SLIDE 6 — Research and references
# =========================================================
set_team(s6)
remove_shape(shape_by_name(s6, "TextBox 8"))
pointer(s6, 0.5, 1.28, 12.3, "Details / Links of the reference and research work")

left = [
    ("The problem in India", [
        ("Standing Committee on Railways (2019). Maintenance of Bridges in Indian Railways – A Review.",
         "prsindia.org/policy/report-summaries/maintenance-of-bridges-in-indian-railways-a-review"),
        ("Gambhira Bridge collapse, Vadodara–Anand, 9 July 2025.",
         "en.wikipedia.org/wiki/Gambhira_Bridge_collapse"),
        ("Railway Board circular 2021/25/CE-III/BR/BMS/Drone: 4D BrIM inspection of railway bridges using UAS.",
         "indianrailways.gov.in → Railway Board → Bridge policy circulars"),
    ]),
    ("Regulation and standards", [
        ("Ministry of Transport, China (Feb 2026). UAV highway-bridge inspection technical guideline, "
         "交办公路〔2026〕8号 (3 m minimum standoff).",
         "xxgk.mot.gov.cn/jigou/glj/202602/P020260228435553861410.pdf"),
        ("Ministry of Civil Aviation. Drone Rules 2021 and type certification.",
         "pib.gov.in/PressReleasePage.aspx?PRID=1833750"),
    ]),
]
right = [
    ("Technical basis", [
        ("Wang et al. RCO-YOLOv5 UAV bridge-defect detection, mAP@0.5 = 91.0%. ASCE J. Perform. Constr. Facil.",
         "doi.org/10.1061/JPCFEV.CFENG-5441"),
        ("Mundt et al. (CVPR 2019). CODEBRIM concrete bridge-defect dataset.",
         "arxiv.org/abs/1904.08486"),
        ("Flotzinger et al. (WACV 2024). dacl10k bridge-damage benchmark.",
         "arxiv.org/abs/2309.00460"),
        ("Milillo et al., Nature Communications (2025): fewer than 1 in 5 long-span bridges are monitored.",
         "NASA Science summary, 12 Dec 2025"),
        ("Flyability Elios 3: LiDAR-SLAM inspection drone for GPS-denied spaces.",
         "flyability.com/elios-3"),
        ("PyTorch torchvision detection models (BSD-3-Clause).",
         "pytorch.org/vision/stable/models.html"),
    ]),
]


def ref_column(slide, x, w, groups, start_num):
    tb = textbox(slide, x, 1.75, w, 5.05)
    tf = tb.text_frame
    n = start_num
    first = True
    for head, items in groups:
        add_para(tf, head, size=14, bold=True, color=BLUE, first=first, space_after=5)
        first = False
        for text, link in items:
            runs = [(f"{n}. ", {"bold": True}), (text + " ", {})]
            if "." in link.split("/")[0] and " " not in link:
                runs.append((link, {"color": BLUE, "link": "https://" + link, "size": 10.5}))
            else:
                runs.append((link, {"color": MUTED, "size": 10.5, "italic": True}))
            add_para(tf, runs, size=12, space_after=9)
            n += 1
        # spacer
        add_para(tf, "", size=8)
    return n


nxt = ref_column(s6, 0.5, 5.95, left, 1)
ref_column(s6, 6.88, 5.95, right, nxt)

prs.save("SIH26201_AVIAN_Idea.pptx")
print("saved")
